from pathlib import Path
from uuid import uuid4
import os
import shutil
import csv

from dotenv import load_dotenv
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings
from langchain_chroma import Chroma
from pypdf import PdfReader
from docx import Document as DocxDocument
from openpyxl import load_workbook


# -------------------------
# 1. Define project paths
# -------------------------

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
CHROMA_DIR = BASE_DIR / "chroma_db"


# -------------------------
# 2. Load environment variables
# -------------------------

load_dotenv(BASE_DIR / ".env")

if not os.getenv("OPENAI_API_KEY"):
    raise ValueError(
        "OPENAI_API_KEY was not found. Add it to your .env file."
    )


# -------------------------
# 3. Document loading
# -------------------------

documents = []


def load_txt(file_path):
    text = file_path.read_text(encoding="utf-8")

    return Document(
        page_content=text,
        metadata={
            "source": str(file_path.relative_to(DATA_DIR)),
            "file_type": file_path.suffix.lower()
        }
    )


def load_pdf(file_path):
    reader = PdfReader(str(file_path))

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    text = "\n\n".join(pages)

    return Document(
        page_content=text,
        metadata={
            "source": str(file_path.relative_to(DATA_DIR)),
            "file_type": ".pdf"
        }
    )


def load_docx(file_path):
    doc = DocxDocument(str(file_path))

    paragraphs = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    text = "\n\n".join(paragraphs)

    return Document(
        page_content=text,
        metadata={
            "source": str(file_path.relative_to(DATA_DIR)),
            "file_type": ".docx"
        }
    )


def load_csv(file_path):
    rows = []

    with open(file_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)

        for row in reader:
            rows.append(" | ".join(row))

    text = "\n".join(rows)

    return Document(
        page_content=text,
        metadata={
            "source": str(file_path.relative_to(DATA_DIR)),
            "file_type": ".csv"
        }
    )


def load_xlsx(file_path):
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True
    )

    sections = []

    for worksheet in workbook.worksheets:

        sections.append(
            f"Worksheet: {worksheet.title}"
        )

        for row in worksheet.iter_rows(values_only=True):

            values = [
                str(value)
                for value in row
                if value is not None
            ]

            if values:
                sections.append(" | ".join(values))

    text = "\n".join(sections)

    return Document(
        page_content=text,
        metadata={
            "source": str(file_path.relative_to(DATA_DIR)),
            "file_type": ".xlsx"
        }
    )


def load_document(file_path):

    extension = file_path.suffix.lower()

    if extension in [".txt", ".md"]:
        return load_txt(file_path)

    if extension == ".pdf":
        return load_pdf(file_path)

    if extension == ".docx":
        return load_docx(file_path)

    if extension == ".csv":
        return load_csv(file_path)

    if extension == ".xlsx":
        return load_xlsx(file_path)

    return None


# -------------------------
# 4. Load all supported files
# -------------------------

supported_extensions = {
    ".txt",
    ".md",
    ".pdf",
    ".docx",
    ".csv",
    ".xlsx"
}

for file_path in sorted(DATA_DIR.rglob("*")):

    if not file_path.is_file():
        continue

    if file_path.suffix.lower() not in supported_extensions:
        continue

    try:

        document = load_document(file_path)

        if document and document.page_content.strip():
            documents.append(document)

            print(
                f"Loaded: "
                f"{document.metadata['source']}"
            )

    except Exception as e:

        print(
            f"ERROR loading {file_path}: {e}"
        )


print()
print(f"Loaded {len(documents)} documents.")


# -------------------------
# 5. Split documents
# -------------------------

text_splitter = RecursiveCharacterTextSplitter(
    chunk_size=500,
    chunk_overlap=75
)

chunks = text_splitter.split_documents(documents)

print(f"Created {len(chunks)} chunks.")


# -------------------------
# 6. Create embeddings
# -------------------------

embeddings = OpenAIEmbeddings(
    model="text-embedding-3-small"
)


# -------------------------
# 7. Rebuild Chroma database
# -------------------------

if CHROMA_DIR.exists():
    shutil.rmtree(CHROMA_DIR)


vector_store = Chroma(
    collection_name="hatfield_helpdesk",
    embedding_function=embeddings,
    persist_directory=str(CHROMA_DIR)
)


# -------------------------
# 8. Store chunks
# -------------------------

ids = [
    str(uuid4())
    for _ in chunks
]

vector_store.add_documents(
    documents=chunks,
    ids=ids
)


# -------------------------
# 9. Completion message
# -------------------------

print()
print("Embeddings created successfully.")
print(f"Stored {len(chunks)} chunks in ChromaDB.")
print(f"Database location: {CHROMA_DIR}")